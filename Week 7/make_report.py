"""Generate the Week 7 model-evaluation report as a DOCX file."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUT = BASE / "report.docx"
FIGURES = BASE / "screenshots"
RESULTS = json.loads((FIGURES / "metrics.json").read_text(encoding="utf-8"))
CUTOFFS = json.loads((FIGURES / "cutoffs.json").read_text(encoding="utf-8"))
GITHUB_URL = "https://github.com/Chief-spartan-117/week-7-ir-evaluation"


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    properties.append(element)


def cell_text(cell, text: str, bold: bool = False, color: str | None = None, size: float = 8.5) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def hyperlink(paragraph, label: str, url: str) -> None:
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_apa_reference(doc: Document, segments: list[tuple[str, bool]], url: str | None = None) -> None:
    """Add an APA 7 reference entry with a 0.5-inch hanging indent."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.space_after = Pt(8)
    for value, italic in segments:
        run = paragraph.add_run(value)
        run.italic = italic
    if url:
        hyperlink(paragraph, url, url)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.25)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (
        ("Title", 28, "16324F"), ("Heading 1", 18, "16324F"),
        ("Heading 2", 13, "2563EB"), ("Heading 3", 11, "334155"),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    caption = doc.styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("475569")
    if "Code" not in doc.styles:
        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code"]
    code.font.name = "Menlo"
    code.font.size = Pt(8.2)
    code.font.color.rgb = RGBColor.from_string("E2E8F0")
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.14)
    properties = code.element.get_or_add_pPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), "172033")
    properties.append(fill)

    header = section.header.paragraphs[0]
    header.text = "WEEK 7  /  RETRIEVAL MODEL EVALUATION"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    add_page_number(section.footer.paragraphs[0])


def caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)


def figure(doc: Document, filename: str, description: str, width: float = 6.65) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    caption(doc, description)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def code(doc: Document, content: str) -> None:
    doc.add_paragraph(content, style="Code")


def title_page(doc: Document) -> None:
    top = doc.add_paragraph("WEEK 7 ASSIGNMENT", style="Subtitle")
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\n")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Evaluation of Information\nRetrieval Models")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Binary Overlap, TF–IDF Cosine, and BM25 on the CISI Project")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string("2563EB")
    doc.add_paragraph("\n")

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details = [
        ("Student", "Sagar Shrestha"),
        ("Student ID", "____________________________"),
        ("Module / Lecturer", "____________________________"),
        ("Submission date", date.today().strftime("%d %B %Y")),
        ("Language", "Python 3"),
    ]
    for row, (label, value) in zip(table.rows, details):
        shade(row.cells[0], "E2E8F0")
        cell_text(row.cells[0], label, True, size=9)
        cell_text(row.cells[1], value, size=9)
    doc.add_paragraph("\n")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Repository: ").bold = True
    hyperlink(paragraph, GITHUB_URL, GITHUB_URL)
    note = doc.add_paragraph("Remote publication requires GitHub authentication before submission.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.italic = True
    note.runs[0].font.color.rgb = RGBColor.from_string("9A3412")
    doc.add_page_break()


def result_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headings = ["Model", "P@5", "R@5", "MAP", "MRR", "nDCG@5", "ms/query"]
    for cell, label in zip(table.rows[0].cells, headings):
        shade(cell, "1E3A5F")
        cell_text(cell, label, True, "FFFFFF", 8.2)
    for name, result in RESULTS["models"].items():
        mean = result["mean"]
        values = [
            name, f'{mean["precision"]:.3f}', f'{mean["recall"]:.3f}',
            f'{mean["ap"]:.3f}', f'{mean["rr"]:.3f}', f'{mean["ndcg"]:.3f}',
            f'{result["mean_latency_ms"]:.4f}',
        ]
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell_text(cell, value, size=7.8)


def per_query_table(doc: Document) -> None:
    models = list(RESULTS["models"])
    queries = [row["query_id"] for row in RESULTS["models"][models[0]]["rows"]]
    table = doc.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ["Model"] + queries):
        shade(cell, "1E3A5F")
        cell_text(cell, label, True, "FFFFFF", 7.5)
    for model in models:
        values = [model] + [f'{row["ap"]:.3f}' for row in RESULTS["models"][model]["rows"]]
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell_text(cell, value, size=7.2)


def cutoff_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ["Model", "R@1", "R@3", "R@5", "R@10"]):
        shade(cell, "1E3A5F")
        cell_text(cell, label, True, "FFFFFF")
    for model in RESULTS["models"]:
        values = [model] + [f'{CUTOFFS[str(k)]["models"][model]["mean"]["recall"]:.3f}' for k in (1, 3, 5, 10)]
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell_text(cell, value)


def build() -> Document:
    doc = Document()
    configure(doc)
    title_page(doc)

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This report evaluates the information retrieval project proposed and implemented in Week 6. The original TF–IDF cosine "
        "system was retained and compared with a binary term-overlap baseline and Okapi BM25. All models received the same documents, "
        "preprocessing, queries, relevance judgments, and output cut-offs. Effectiveness was measured with Precision@k, Recall@k, "
        "Mean Average Precision (MAP), Mean Reciprocal Rank (MRR), and normalized Discounted Cumulative Gain (nDCG@k). Mean query "
        "latency was also recorded as a descriptive efficiency measure."
    )
    doc.add_paragraph(
        "On the bundled offline demonstration, all models achieved R@5 = 1.000 and P@5 = 0.325. TF–IDF produced the highest MAP "
        "(0.948), binary overlap produced the highest MRR (1.000), and binary overlap also had the highest nDCG@5 (0.970). BM25 "
        "obtained MAP = 0.927 and nDCG@5 = 0.952. Differences came from ordering rather than failure to retrieve relevant records. "
        "Because only 18 documents and eight manually judged queries were available offline, these results demonstrate the evaluation "
        "implementation and must not be treated as full CISI benchmark results."
    )
    doc.add_heading("Contents", level=1)
    for item in [
        "1. Introduction", "2. Evaluation objectives and questions", "3. Dataset and experimental controls",
        "4. Retrieval models", "5. Evaluation metrics", "6. Implementation",
        "7. Results", "8. Detailed analysis", "9. Verification and validity",
        "10. Reproducibility and GitHub", "11. Conclusion", "References", "Appendix A — Project structure",
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Week 6 established a basic CISI search engine using an inverted index, shared text preprocessing, logarithmic TF–IDF weights, "
        "and cosine similarity. Week 7 moves from implementation to comparative evaluation. A retrieval score is useful only when its "
        "ranking can be compared with known relevance judgments. The evaluation therefore treats every query as a test case and measures "
        "both how many relevant documents are returned and how early they appear."
    )
    doc.add_paragraph(
        "The comparison includes a deliberately simple baseline, the Week 6 model, and a widely used probabilistic ranking function. "
        "This design makes the value of weighting and length normalization observable while keeping the experiment small enough to audit."
    )

    doc.add_heading("2. Evaluation Objectives and Questions", level=1)
    doc.add_heading("2.1 Objectives", level=2)
    bullets(doc, [
        "Implement precision, recall, AP/MAP, RR/MRR, and nDCG without relying on a black-box evaluation library.",
        "Compare three retrieval models under identical experimental conditions.",
        "Inspect aggregate scores and per-query variation rather than declaring a winner from one number.",
        "Measure effectiveness at multiple result cut-offs (k = 1, 3, 5, and 10).",
        "Verify metric formulas and retrieval behavior with automated tests.",
        "Produce reproducible code, inline comments, screenshots, analysis, and a DOCX report.",
    ])
    doc.add_heading("2.2 Research questions", level=2)
    numbered(doc, [
        "Do term weighting and document-length normalization improve the ranking over binary overlap?",
        "Do the models differ in first-result quality, overall ranking quality, or only their raw score scales?",
        "How does increasing k affect precision and recall?",
        "Are apparent performance differences credible given the size and construction of the offline sample?",
    ])
    doc.add_heading("2.3 Non-goals", level=2)
    doc.add_paragraph(
        "The experiment does not tune parameters on the test queries, train a learning-to-rank model, introduce embeddings, or claim "
        "statistical significance. Those additions would complicate the baseline without improving the assignment’s central evidence."
    )

    doc.add_heading("3. Dataset and Experimental Controls", level=1)
    doc.add_heading("3.1 CISI project", level=2)
    doc.add_paragraph(
        "The selected project is the CISI information-retrieval collection used in Week 6. The University of Glasgow resource lists "
        "1,460 documents and 112 queries and supplies separate document, query, and relevance-assessment files. A standard IR test "
        "collection requires exactly these three elements: a corpus, information needs, and relevance judgments (Manning et al., "
        "2008; University of Glasgow Information Retrieval Group, n.d.-a, n.d.-b)."
    )
    doc.add_heading("3.2 Offline execution layer", level=2)
    doc.add_paragraph(
        "Direct external downloads are unavailable in the execution environment, so the included experiment uses the same 18-record, "
        "eight-query demonstration layer prepared for Week 6. Each query has one or two binary relevance judgments. The repository also "
        "includes a downloader and parsers for the complete CISI files, so the identical evaluation can be rerun on the full collection."
    )
    doc.add_heading("3.3 Controlled comparison", level=2)
    bullets(doc, [
        "Same lowercasing, Unicode normalization, tokenization, stop list, and suffix stemming for all models.",
        "Same title repetition (a modest field boost) and searchable document fields.",
        "Same eight queries and binary relevance judgments.",
        "Same top-k values and full candidate ranking for AP and RR.",
        "Default model parameters fixed before evaluation: BM25 k1 = 1.5 and b = 0.75.",
        "Deterministic document-ID tie breaking, making every run reproducible.",
    ])
    figure(doc, "evaluation_architecture.png", "Figure 1. Controlled pipeline: only the scoring model changes.")
    doc.add_page_break()

    doc.add_heading("4. Retrieval Models", level=1)
    doc.add_heading("4.1 Binary term overlap", level=2)
    doc.add_paragraph(
        "The baseline converts the query to a set of unique terms and scores each document by the proportion of query terms it contains. "
        "It ignores term frequency, inverse document frequency, and document length. Its simplicity establishes whether more sophisticated "
        "weighting adds value. Equal scores are resolved by document ID."
    )
    code(doc, "score_overlap(q,d) = |unique(q) ∩ terms(d)| / |unique(q)|")
    doc.add_heading("4.2 TF–IDF cosine", level=2)
    doc.add_paragraph(
        "This is the Week 6 model. Logarithmic term frequency limits the influence of repetition, IDF favors discriminating terms, and "
        "cosine normalization controls vector length. The query uses the same term weighting as the documents (Manning et al., 2008)."
    )
    code(doc, "tfw(t,d) = 1 + ln(tf(t,d))\nidf(t) = ln((N + 1) / (df(t) + 1)) + 1\nscore(q,d) = (q · d) / (||q||₂ ||d||₂)")
    doc.add_heading("4.3 Okapi BM25", level=2)
    doc.add_paragraph(
        "BM25 combines probabilistic IDF with saturating term-frequency evidence and explicit document-length normalization. Parameter "
        "k1 controls term-frequency saturation; b determines the strength of length normalization. Conventional baseline values were "
        "used without tuning to the demonstration queries (Robertson & Zaragoza, 2009)."
    )
    code(doc, "idf(t) = ln(1 + (N - df(t) + 0.5) / (df(t) + 0.5))\nscore(q,d) = Σ idf(t) × tf(t,d)(k1+1) / [tf(t,d)+k1(1-b+b|d|/avgdl)]")
    figure(doc, "code_bm25.png", "Figure 2. Commented BM25 implementation used in the experiment.", width=6.55)

    doc.add_heading("5. Evaluation Metrics", level=1)
    doc.add_heading("5.1 Precision and recall at k", level=2)
    code(doc, "P@k = relevant documents in top k / k\nR@k = relevant documents in top k / all relevant documents")
    doc.add_paragraph(
        "Precision measures result purity; recall measures coverage. They often move in opposite directions as k grows. In this sample, "
        "P@5 has a low ceiling because each query has at most two judged relevant documents."
    )
    doc.add_heading("5.2 Average Precision and MAP", level=2)
    doc.add_paragraph(
        "Average Precision averages the precision observed whenever a relevant document is encountered, dividing by the total number "
        "of known relevant documents. Missing relevant documents therefore contribute zero. MAP is the arithmetic mean of AP over queries."
    )
    code(doc, "AP(q) = [Σ Precision@r × rel(r)] / number_relevant(q)\nMAP = mean(AP(q))")
    doc.add_heading("5.3 Reciprocal rank and MRR", level=2)
    doc.add_paragraph(
        "Reciprocal Rank focuses only on the position of the first relevant result. It is appropriate when users primarily need one "
        "satisfactory answer. MRR averages this value over queries."
    )
    code(doc, "RR(q) = 1 / rank_of_first_relevant_document\nMRR = mean(RR(q))")
    doc.add_heading("5.4 nDCG at k", level=2)
    doc.add_paragraph(
        "Discounted Cumulative Gain rewards relevant documents more when they occur near the top. Dividing by the ideal DCG makes the "
        "score comparable across queries with different numbers of relevant documents. With binary judgments, every relevant record "
        "contributes a gain of one (Järvelin & Kekäläinen, 2002)."
    )
    code(doc, "DCG@k = Σ rel(r) / log₂(r + 1)\nnDCG@k = DCG@k / IDCG@k")
    doc.add_heading("5.5 Latency", level=2)
    doc.add_paragraph(
        "Mean query latency is measured with a monotonic high-resolution timer over 500 repetitions of all eight queries. It describes "
        "this implementation on this machine; it is not a general speed benchmark because the collection is tiny and times are below "
        "one millisecond."
    )
    doc.add_page_break()

    doc.add_heading("6. Implementation", level=1)
    doc.add_heading("6.1 Metric calculation", level=2)
    figure(doc, "code_metrics.png", "Figure 3. Actual metric function with validation and inline comments.", width=6.55)
    doc.add_paragraph(
        "The implementation first computes P@k and R@k from the cut-off list. It then scans the full ranked list once to accumulate AP, "
        "find the first relevant rank, and calculate DCG up to k. Ideal DCG is constructed from the smaller of k and the number of "
        "relevant documents. Queries without judgments are excluded before aggregation; a judged query with an empty relevance set returns zeros."
    )
    doc.add_heading("6.2 Model evaluation loop", level=2)
    numbered(doc, [
        "Create shared corpus statistics and instantiate the three models.",
        "For each judged query, retrieve a complete non-zero ranking.",
        "Calculate the five effectiveness measures at the selected k.",
        "Repeat top-k retrieval for descriptive latency measurement.",
        "Average each metric across queries and serialize all rows to JSON.",
    ])
    doc.add_heading("6.3 Files and responsibilities", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["File", "Responsibility"]):
        shade(cell, "1E3A5F"); cell_text(cell, label, True, "FFFFFF")
    for filename, purpose in [
        ("retrieval_models.py", "Preprocessing, corpus statistics, overlap, TF–IDF, and BM25"),
        ("evaluate_models.py", "Parsers, metric formulas, aggregation, timing, CLI, JSON output"),
        ("download_data.py", "Full CISI download and structural validation"),
        ("demo.py", "Side-by-side top-three ranking example"),
        ("tests/", "Known-answer metric tests and cross-model retrieval tests"),
        ("make_screenshots.py", "Regenerates figures from current code and output"),
        ("make_report.py", "Builds this report from JSON metrics and PNG evidence"),
    ]:
        cells = table.add_row().cells
        cell_text(cells[0], filename)
        cell_text(cells[1], purpose)

    doc.add_heading("7. Results", level=1)
    doc.add_heading("7.1 Aggregate results at k = 5", level=2)
    result_table(doc)
    caption(doc, "Table 1. Aggregate offline demonstration results. Higher is better except latency.")
    figure(doc, "model_comparison.png", "Figure 4. Aggregate model comparison at k = 5.", width=6.45)
    figure(doc, "evaluation_output.png", "Figure 5. Actual terminal output used to populate Table 1.", width=6.55)

    doc.add_heading("7.2 Per-query Average Precision", level=2)
    per_query_table(doc)
    caption(doc, "Table 2. AP for every query and model.")
    figure(doc, "per_query_ap.png", "Figure 6. Per-query AP heat map; variation is concentrated in Q3, Q4, and Q8.", width=6.45)

    doc.add_heading("7.3 Recall at multiple cut-offs", level=2)
    cutoff_table(doc)
    caption(doc, "Table 3. Mean recall as k increases.")
    figure(doc, "recall_cutoffs.png", "Figure 7. All models reach complete judged recall by rank 3.", width=6.45)

    doc.add_heading("7.4 Ranking example", level=2)
    figure(doc, "ranking_output.png", "Figure 8. Same query, different ranking functions and score scales.", width=6.55)
    doc.add_paragraph(
        "For Q8, binary overlap ranks S012 first and therefore obtains RR = 1.0. TF–IDF and BM25 place the non-relevant S003 first "
        "because it matches several weighted bibliographic terms, pushing relevant records to ranks two and three. The example also "
        "shows why raw scores cannot be compared across models: overlap is bounded by one, cosine is normalized, and BM25 is unbounded."
    )
    doc.add_page_break()

    doc.add_heading("8. Detailed Analysis", level=1)
    doc.add_heading("8.1 Answer to research question 1", level=2)
    doc.add_paragraph(
        "Term weighting did not uniformly outperform the baseline. TF–IDF achieved the highest MAP (0.948), improving on overlap by "
        "0.010 absolute, because it ranked both relevant items ideally for Q3 and Q4. However, binary overlap performed better on Q8 "
        "and achieved the best MRR and nDCG@5. BM25 did not beat TF–IDF in this sample. The evidence therefore supports no universal "
        "winner; it shows that the effect of weighting depends on the query and corpus."
    )
    doc.add_heading("8.2 First-result quality versus complete ranking", level=2)
    doc.add_paragraph(
        "Binary overlap returned a relevant first result for every query, giving P@1 = 1.000 and MRR = 1.000. TF–IDF and BM25 each "
        "missed at rank one for Q8, giving P@1 = 0.875 and MRR = 0.938. Yet TF–IDF’s higher MAP shows that after the first position it "
        "often ordered multiple relevant records better. MRR and MAP answer different user needs and should be reported together."
    )
    doc.add_heading("8.3 Effect of k", level=2)
    doc.add_paragraph(
        "Mean recall rises from 0.688 for overlap and 0.625 for the weighted models at k = 1 to 1.000 for all models at k = 3. "
        "Recall remains 1.000 thereafter. Precision necessarily falls as k increases: the mean is 0.325 at k = 5 and 0.162 at k = 10 "
        "because the judgment sets contain only one or two relevant documents. This is the expected precision–recall trade-off rather "
        "than evidence that the models become worse."
    )
    doc.add_heading("8.4 Per-query stability", level=2)
    doc.add_paragraph(
        "Five of eight queries produce AP = 1.000 for every model. The aggregate differences are driven by only Q3, Q4, and Q8. "
        "TF–IDF beats overlap on two of those queries; overlap beats TF–IDF on one; the remaining five are ties. Such concentration means "
        "a change to a single query could change the model ordering. Reporting only a mean would conceal this instability."
    )
    doc.add_heading("8.5 Efficiency", level=2)
    doc.add_paragraph(
        "Binary overlap is the fastest model in the recorded run because it performs only set-style accumulation. BM25 adds length "
        "normalization and term saturation; TF–IDF adds query and document vector normalization. All recorded times are around one "
        "hundredth of a millisecond per query, so operating-system and interpreter noise are large relative to the work. The results "
        "should be read as an implementation sanity check, not as a basis for production capacity planning."
    )
    doc.add_heading("8.6 Recommendation", level=2)
    doc.add_paragraph(
        "TF–IDF remains the recommended Week 6 model because it has the best MAP, is easy to explain, and requires no parameter tuning. "
        "Binary overlap should remain as a regression baseline. BM25 should be reevaluated on the full CISI collection before rejection: "
        "its document-length normalization has little opportunity to help in an 18-record demonstration corpus, and its parameters were not tuned."
    )

    doc.add_heading("9. Verification and Validity", level=1)
    doc.add_heading("9.1 Automated tests", level=2)
    doc.add_paragraph(
        "Seven tests passed. Four metric tests use rankings with manually calculable answers, including a relevant document missing from "
        "the returned list. Three parameterized model tests verify that all systems rank an obvious match first, return an empty list "
        "for unknown terms, and reject an invalid top-k value."
    )
    figure(doc, "tests_output.png", "Figure 9. Seven automated tests completed successfully.", width=6.55)
    doc.add_heading("9.2 Threats to validity", level=2)
    bullets(doc, [
        "External validity: 18 demonstration documents cannot represent performance on all 1,460 CISI records.",
        "Judgment validity: demonstration relevance sets were manually prepared and are not official CISI assessments.",
        "Statistical validity: eight queries are too few for reliable significance testing or broad model claims.",
        "Tie sensitivity: binary overlap creates many equal scores, so deterministic document-ID order can help or hurt a query.",
        "Parameter validity: BM25 k1 and b are conventional defaults, not values selected on development data.",
        "Timing validity: sub-millisecond measurements on a tiny corpus are dominated by local runtime conditions.",
    ])
    doc.add_heading("9.3 Mitigations", level=2)
    doc.add_paragraph(
        "The report labels every demonstration result, publishes per-query values, retains fixed controls, uses known-answer tests, and "
        "provides a full-CISI execution path. No statistical or benchmark claim is made from the offline sample."
    )

    doc.add_heading("10. Reproducibility and GitHub", level=1)
    doc.add_paragraph("Reproduce the submitted offline results:")
    code(doc, "python3 evaluate_models.py --sample --top-k 5 --json results.json\npython3 demo.py\npython3 -m unittest discover -s tests -v")
    doc.add_paragraph("Evaluate the complete public collection when network access is available:")
    code(doc, "python3 download_data.py\npython3 evaluate_models.py --data-dir data --top-k 10 --json full_cisi_results.json")
    doc.add_paragraph("Regenerate every snapshot and this DOCX report:")
    code(doc, "python3 make_screenshots.py\npython3 make_report.py")
    doc.add_paragraph("Publish the prepared repository:")
    code(doc, "gh auth login\ngh repo create Chief-spartan-117/week-7-ir-evaluation --public\ngit push -u origin main")
    p = doc.add_paragraph()
    p.add_run("Intended GitHub link: ").bold = True
    hyperlink(p, GITHUB_URL, GITHUB_URL)
    doc.add_paragraph(
        "Before submission, open the link in a private browser window to confirm public access. The local repository already uses the "
        "main branch and has this URL configured as origin."
    )

    doc.add_heading("11. Conclusion", level=1)
    doc.add_paragraph(
        "Week 7 adds a reproducible evaluation layer to the Week 6 information-retrieval project. The work implements five standard "
        "effectiveness measures, compares three models under controlled conditions, records multi-cutoff behavior, tests the metric "
        "formulas, and preserves per-query evidence. TF–IDF has the best MAP on the demonstration corpus, while binary overlap has the "
        "best first-result performance and nDCG@5. These differences show why no single metric is sufficient. The most important next "
        "step is to run the same scripts on all official CISI queries and judgments; only that experiment can support a defensible "
        "model-selection claim."
    )

    doc.add_page_break()
    doc.add_heading("References", level=1)
    add_apa_reference(doc, [
        ("GianRomani. (n.d.). ", False),
        ("CISI-project-MLOps", True),
        (" [Data set]. GitHub. Retrieved October 12, 2026, from ", False),
    ], "https://github.com/GianRomani/CISI-project-MLOps")
    add_apa_reference(doc, [
        ("Järvelin, K., & Kekäläinen, J. (2002). Cumulated gain-based evaluation of IR techniques. ", False),
        ("ACM Transactions on Information Systems, 20", True),
        ("(4), 422–446. ", False),
    ], "https://doi.org/10.1145/582415.582418")
    add_apa_reference(doc, [
        ("Manning, C. D., Raghavan, P., & Schütze, H. (2008). ", False),
        ("Introduction to information retrieval", True),
        (". Cambridge University Press. ", False),
    ], "https://nlp.stanford.edu/IR-book/")
    add_apa_reference(doc, [
        ("Robertson, S., & Zaragoza, H. (2009). The probabilistic relevance framework: BM25 and beyond. ", False),
        ("Foundations and Trends in Information Retrieval, 3", True),
        ("(4), 333–389. ", False),
    ], "https://doi.org/10.1561/1500000019")
    add_apa_reference(doc, [
        ("University of Glasgow Information Retrieval Group. (n.d.-a). ", False),
        ("CISI collection", True),
        (". Retrieved October 12, 2026, from ", False),
    ], "https://ir.dcs.gla.ac.uk/resources/test_collections/cisi/")
    add_apa_reference(doc, [
        ("University of Glasgow Information Retrieval Group. (n.d.-b). ", False),
        ("Test collections", True),
        (". Retrieved October 12, 2026, from ", False),
    ], "https://ir.dcs.gla.ac.uk/resources/test_collections/")

    doc.add_heading("Appendix A — Project Structure", level=1)
    code(doc, """Week 7/
├── data/
│   ├── sample_documents.json
│   └── sample_queries.json
├── screenshots/
├── tests/test_models_and_metrics.py
├── retrieval_models.py
├── evaluate_models.py
├── download_data.py
├── demo.py
├── make_screenshots.py
├── make_report.py
├── README.md
├── requirements.txt
└── report.docx""")
    doc.add_paragraph(
        "The JSON result files are retained with the figures. The report generator reads those files directly, preventing values in the "
        "written tables from drifting away from the executed experiment."
    )
    return doc


if __name__ == "__main__":
    report = build()
    report.save(OUT)
    print(f"Saved {OUT}")

