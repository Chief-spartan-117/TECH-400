"""Generate a polished PowerPoint summary of the Week 6–7 IR project."""

from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
WEEK6 = ROOT / "Week 6"
WEEK7 = ROOT / "Week 7"
OUTPUT = BASE / "Week_8_Project_Findings.pptx"
METRICS = json.loads((WEEK7 / "screenshots/metrics.json").read_text(encoding="utf-8"))
CUTOFFS = json.loads((WEEK7 / "screenshots/cutoffs.json").read_text(encoding="utf-8"))

SLIDE_W = 13.333
SLIDE_H = 7.5

# Presentation palette
NAVY = "0B1020"
NAVY_2 = "151D33"
SLATE = "27344E"
MUTED = "94A3B8"
WHITE = "F8FAFC"
PAPER = "F3F6FB"
INK = "14213D"
BLUE = "4F7CFF"
CYAN = "34D1BF"
ORANGE = "FFB454"
MAGENTA = "EB5CA8"
GREEN = "55C38E"
RED = "F07178"
LINE = "D8E1EF"

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, radius=True, line=None, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x1), Inches(y1), Inches(max(0.01, x2 - x1)), Inches(max(0.01, y2 - y1)),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def text_box(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.04, fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    if fit:
        frame.fit_text(font_family=font, max_size=size)
    return box


def rich_text(slide, runs, x, y, w, h, size=18, color=INK,
              align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for value, run_color, bold in runs:
        run = paragraph.add_run()
        run.text = value
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(run_color or color)
    return box


def bullet_list(slide, items, x, y, w, h, size=17, color=INK,
                bullet_color=BLUE, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.level = 0
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.05
        bullet = paragraph.add_run()
        bullet.text = "●  "
        bullet.font.name = FONT
        bullet.font.size = Pt(size - 3)
        bullet.font.color.rgb = rgb(bullet_color)
        run = paragraph.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
    return box


def add_background(slide, color=PAPER):
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = rgb(color)


def add_header(slide, number: int, kicker: str, title: str, dark=False, source: str | None = None):
    main = WHITE if dark else INK
    muted = MUTED if dark else "64748B"
    text_box(slide, kicker.upper(), 0.72, 0.30, 8.8, 0.26, 9.5, BLUE, True, FONT, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, title, 0.72, 0.62, 11.4, 0.62, 26, main, True, FONT_DISPLAY, valign=MSO_ANCHOR.MIDDLE)
    rect(slide, 12.25, 0.34, 0.48, 0.36, BLUE if dark else NAVY, True)
    text_box(slide, f"{number:02d}", 12.25, 0.34, 0.48, 0.36, 10, WHITE, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    if source:
        text_box(slide, source, 0.72, 7.16, 11.9, 0.17, 7.5, muted, False, FONT, valign=MSO_ANCHOR.MIDDLE)


def add_picture_cover(slide, path: Path, x, y, w, h):
    with Image.open(path) as image:
        img_w, img_h = image.size
    image_ratio = img_w / img_h
    frame_ratio = w / h
    if image_ratio > frame_ratio:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))
        excess = pic.width - Inches(w)
        crop_fraction = excess / (2 * pic.width)
        pic.crop_left = crop_fraction
        pic.crop_right = crop_fraction
        pic.left = Inches(x)
        pic.width = Inches(w)
    else:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        excess = pic.height - Inches(h)
        crop_fraction = excess / (2 * pic.height)
        pic.crop_top = crop_fraction
        pic.crop_bottom = crop_fraction
        pic.top = Inches(y)
        pic.height = Inches(h)
    return pic


def add_card(slide, x, y, w, h, title, body, accent=BLUE, dark=False, number=None):
    fill = NAVY_2 if dark else WHITE
    title_color = WHITE if dark else INK
    body_color = "C9D4E7" if dark else "52637A"
    rect(slide, x, y, w, h, fill, True, line=SLATE if dark else LINE)
    rect(slide, x, y, 0.08, h, accent, False)
    if number is not None:
        rect(slide, x + 0.25, y + 0.25, 0.48, 0.44, accent, True)
        text_box(slide, str(number), x + 0.25, y + 0.25, 0.48, 0.44, 11, NAVY, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        title_x = x + 0.88
        title_w = w - 1.08
    else:
        title_x = x + 0.27
        title_w = w - 0.50
    text_box(slide, title, title_x, y + 0.21, title_w, 0.44, 16, title_color, True, FONT_DISPLAY,
             valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, body, x + 0.27, y + 0.76, w - 0.50, h - 0.94, 11.5, body_color,
             valign=MSO_ANCHOR.TOP)


def add_metric(slide, x, y, w, h, value, label, accent=BLUE, note=None, dark=False):
    fill = NAVY_2 if dark else WHITE
    main = WHITE if dark else INK
    muted = MUTED if dark else "64748B"
    rect(slide, x, y, w, h, fill, True, line=SLATE if dark else LINE)
    text_box(slide, value, x + 0.24, y + 0.18, w - 0.48, 0.56, 27, accent, True, FONT_DISPLAY,
             valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, label, x + 0.24, y + 0.80, w - 0.48, 0.28, 11, main, True,
             valign=MSO_ANCHOR.MIDDLE)
    if note:
        text_box(slide, note, x + 0.24, y + 1.10, w - 0.48, h - 1.25, 8.5, muted,
                 valign=MSO_ANCHOR.TOP)


def add_chevron(slide, x, y, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(0.34), Inches(0.42))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_model_result_row(slide, y, model, color, values):
    text_box(slide, model, 0.92, y, 2.10, 0.34, 13, INK, True, valign=MSO_ANCHOR.MIDDLE)
    keys = [("MAP", values["ap"]), ("MRR", values["rr"]), ("nDCG@5", values["ndcg"])]
    starts = [3.0, 6.2, 9.4]
    for (label, value), start in zip(keys, starts):
        text_box(slide, label, start, y - 0.24, 1.0, 0.22, 8.5, "64748B", True)
        rect(slide, start, y + 0.07, 2.45, 0.22, "E2E8F0", True)
        rect(slide, start, y + 0.07, 2.45 * value, 0.22, color, True)
        text_box(slide, f"{value:.3f}", start + 2.55, y - 0.01, 0.65, 0.32, 10, INK, True,
                 align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)


def set_slide_transition(slide):
    transition = parse_xml(f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>')
    slide._element.insert(-1, transition)


def build_deck() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]

    # 1 — Title
    slide = presentation.slides.add_slide(blank)
    add_background(slide, NAVY)
    rect(slide, 8.75, -0.5, 5.1, 8.3, NAVY_2, True)
    rect(slide, 9.35, 0.55, 3.35, 0.10, BLUE, False)
    rect(slide, 9.35, 0.88, 2.25, 0.10, CYAN, False)
    rect(slide, 9.35, 1.21, 1.25, 0.10, ORANGE, False)
    text_box(slide, "INFORMATION RETRIEVAL · WEEKS 6–7", 0.78, 0.67, 7.5, 0.32, 11, CYAN, True,
             valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "From search system\nto measured evidence", 0.78, 1.25, 7.75, 1.65, 34, WHITE, True,
             FONT_DISPLAY, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "A concise evaluation of Binary Overlap, TF–IDF Cosine, and BM25 on the CISI information-retrieval project.",
             0.82, 3.25, 7.05, 1.00, 17, "C9D4E7", False, valign=MSO_ANCHOR.TOP)
    rect(slide, 0.82, 4.65, 6.85, 0.02, SLATE, False)
    text_box(slide, "Sagar Shrestha", 0.82, 4.98, 3.4, 0.38, 15, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Week 8 Project Findings", 0.82, 5.42, 3.8, 0.30, 10.5, MUTED, False, valign=MSO_ANCHOR.MIDDLE)
    # Right-side evidence stack
    add_metric(slide, 9.35, 2.15, 2.75, 1.38, "1,460", "FULL CISI DOCUMENTS", BLUE, dark=True)
    add_metric(slide, 9.35, 3.75, 2.75, 1.38, "0.948", "BEST DEMO MAP", CYAN, dark=True)
    add_metric(slide, 9.35, 5.35, 2.75, 1.38, "7/7", "TESTS PASSED", ORANGE, dark=True)
    text_box(slide, "01", 12.25, 7.05, 0.45, 0.20, 8, MUTED, True, align=PP_ALIGN.RIGHT)
    set_slide_transition(slide)

    # 2 — Executive snapshot
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 2, "Executive snapshot", "What the project established")
    rich_text(slide, [("A working lexical search system", INK, True),
                      (" was designed in Week 6; Week 7 added controlled model comparison and ranked-retrieval evaluation.", "52637A", False)],
              0.78, 1.43, 11.5, 0.72, 17)
    add_card(slide, 0.78, 2.35, 3.80, 2.45, "Build", "Parsed CISI-style records, normalized text, built an inverted index, and returned deterministic ranked results.", BLUE, number="01")
    add_card(slide, 4.77, 2.35, 3.80, 2.45, "Compare", "Held data and preprocessing constant while changing only the scoring model: overlap, TF–IDF, or BM25.", MAGENTA, number="02")
    add_card(slide, 8.76, 2.35, 3.80, 2.45, "Evaluate", "Measured P@k, R@k, MAP, MRR, nDCG@k, latency, and query-level variation with tested formulas.", CYAN, number="03")
    rect(slide, 0.78, 5.25, 11.78, 1.12, NAVY, True)
    text_box(slide, "Decision", 1.08, 5.50, 1.10, 0.30, 10, ORANGE, True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Keep TF–IDF as the explainable baseline; validate BM25 on the full CISI collection before selecting a production model.",
             2.20, 5.40, 9.85, 0.50, 16, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    set_slide_transition(slide)

    # 3 — Problem
    slide = presentation.slides.add_slide(blank)
    add_background(slide, NAVY)
    add_header(slide, 3, "Problem statement", "How can relevant documents be ranked—and trusted?", dark=True)
    text_box(slide, "Natural-language queries rarely match document wording exactly. A useful system must rank partial matches, then prove that the ranking is effective.",
             0.78, 1.52, 6.05, 1.20, 21, WHITE, True, FONT_DISPLAY)
    text_box(slide, "The project therefore had two linked problems:", 0.78, 3.00, 5.7, 0.36, 12, MUTED, True)
    add_card(slide, 0.78, 3.52, 5.75, 1.18, "Retrieval problem", "Transform documents and queries into comparable representations, then produce a ranked top-k list.", BLUE, dark=True, number="A")
    add_card(slide, 0.78, 4.90, 5.75, 1.18, "Evaluation problem", "Use relevance judgments and complementary metrics to determine whether early and complete retrieval are good.", CYAN, dark=True, number="B")
    # Right visual: question → evidence
    rect(slide, 7.45, 1.55, 4.95, 4.92, NAVY_2, True, line=SLATE)
    text_box(slide, "USER NEED", 7.90, 1.95, 1.7, 0.30, 9, BLUE, True)
    text_box(slide, '“automatic indexing and document retrieval”', 7.90, 2.35, 3.95, 0.72, 19, WHITE, True, FONT_DISPLAY)
    line(slide, 7.90, 3.34, 11.80, 3.36, SLATE)
    for index, (rank, label, color) in enumerate([
        ("01", "Automatic Indexing Experiments", CYAN),
        ("02", "Abstracting Concepts and Methods", BLUE),
        ("03", "Construction and Use of a Thesaurus", ORANGE),
    ]):
        y = 3.62 + index * 0.72
        rect(slide, 7.90, y, 0.54, 0.45, color, True)
        text_box(slide, rank, 7.90, y, 0.54, 0.45, 10, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        text_box(slide, label, 8.65, y, 3.05, 0.45, 11.5, WHITE, index == 0, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "The ordering—not just the matches—is the product.", 7.90, 5.94, 3.95, 0.26, 9.5, MUTED, False)
    set_slide_transition(slide)

    # 4 — Dataset
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 4, "Dataset", "CISI: a compact, relevance-judged IR collection",
               source="Source: University of Glasgow IR Test Collections · ir.dcs.gla.ac.uk/resources/test_collections/cisi/")
    add_metric(slide, 0.78, 1.48, 3.65, 1.60, "1,460", "DOCUMENTS", BLUE, "Library and information-science records")
    add_metric(slide, 4.64, 1.48, 3.65, 1.60, "112", "QUERIES", MAGENTA, "Natural-language information needs")
    add_metric(slide, 8.50, 1.48, 3.65, 1.60, "3 files", "TEST COLLECTION", CYAN, "Documents · queries · relevance judgments")
    text_box(slide, "Two evidence layers", 0.78, 3.52, 3.0, 0.36, 16, INK, True, FONT_DISPLAY)
    add_card(slide, 0.78, 4.05, 5.55, 1.82, "Full benchmark target", "CISI.ALL, CISI.QRY, and CISI.REL support evaluation over the official 1,460-document collection. The code includes download, parsing, and validation.", BLUE, number="FULL")
    add_card(slide, 6.55, 4.05, 5.60, 1.82, "Executed offline demonstration", "18 concise records, 8 queries, and manual binary judgments prove the pipeline in a restricted environment. These scores are illustrative—not benchmark claims.", ORANGE, number="DEMO")
    rect(slide, 0.78, 6.20, 11.37, 0.52, "FFF3DF", True, line="F6D69A")
    text_box(slide, "Interpretation rule  →  never generalize demonstration scores to the full CISI collection.",
             1.02, 6.30, 10.85, 0.26, 11, "8A4B08", True, valign=MSO_ANCHOR.MIDDLE)
    set_slide_transition(slide)

    # 5 — Methodology
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 5, "Methodology", "One pipeline, three scoring models")
    stages = [
        ("01", "Parse", ".I / .T / .A / .W", BLUE),
        ("02", "Normalize", "lowercase · tokenize", CYAN),
        ("03", "Filter", "stop words · stemming", GREEN),
        ("04", "Index", "term → doc / tf", ORANGE),
        ("05", "Score", "model-specific", MAGENTA),
        ("06", "Rank", "top-k results", BLUE),
    ]
    for index, (number, title, detail, color) in enumerate(stages):
        x = 0.68 + index * 2.10
        rect(slide, x, 1.62, 1.78, 1.52, WHITE, True, line=LINE)
        rect(slide, x + 0.18, 1.80, 0.48, 0.40, color, True)
        text_box(slide, number, x + 0.18, 1.80, 0.48, 0.40, 9, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        text_box(slide, title, x + 0.18, 2.28, 1.42, 0.30, 14, INK, True, FONT_DISPLAY)
        text_box(slide, detail, x + 0.18, 2.66, 1.42, 0.28, 8.5, "64748B")
        if index < len(stages) - 1:
            add_chevron(slide, x + 1.82, 2.18, color)
    text_box(slide, "Controlled comparison", 0.78, 3.70, 3.0, 0.35, 16, INK, True, FONT_DISPLAY)
    rect(slide, 0.78, 4.22, 11.75, 1.78, NAVY, True)
    controls = [
        ("Same data", "Documents, queries, judgments"),
        ("Same representation", "Title boost and preprocessing"),
        ("Same cut-offs", "k = 1, 3, 5, 10"),
        ("Only score changes", "Fair model comparison"),
    ]
    for index, (title, body) in enumerate(controls):
        x = 1.05 + index * 2.85
        text_box(slide, title, x, 4.61, 2.25, 0.28, 12, CYAN if index < 3 else ORANGE, True)
        text_box(slide, body, x, 5.02, 2.28, 0.48, 10, "C9D4E7")
        if index < 3:
            line(slide, x + 2.43, 4.61, x + 2.45, 5.55, SLATE)
    text_box(slide, "Design principle: hold everything constant except the retrieval function.", 0.78, 6.38, 11.75, 0.35, 12, "52637A", True, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # 6 — Models
    slide = presentation.slides.add_slide(blank)
    add_background(slide, NAVY)
    add_header(slide, 6, "Retrieval models", "Three increasingly expressive lexical baselines", dark=True,
               source="Methods: Manning, Raghavan & Schütze (2008); Robertson & Zaragoza (2009)")
    model_cards = [
        (0.78, "Binary overlap", "TERM PRESENCE", "Coverage of unique query terms", "Fast and transparent", "No rarity or length signal", BLUE),
        (4.55, "TF–IDF cosine", "WEIGHTED VECTOR", "log TF × smooth IDF × cosine", "Best demonstration MAP", "Lexical mismatch remains", MAGENTA),
        (8.32, "BM25", "PROBABILISTIC", "IDF + TF saturation + length", "Strong classic baseline", "Parameters need validation", CYAN),
    ]
    for x, title, label, formula, strength, tradeoff, color in model_cards:
        rect(slide, x, 1.58, 3.50, 4.78, NAVY_2, True, line=SLATE)
        rect(slide, x + 0.28, 1.88, 1.45, 0.30, color, True)
        text_box(slide, label, x + 0.28, 1.88, 1.45, 0.30, 7.5, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        text_box(slide, title, x + 0.28, 2.42, 2.92, 0.48, 20, WHITE, True, FONT_DISPLAY)
        text_box(slide, formula, x + 0.28, 3.12, 2.92, 0.70, 12, "C9D4E7", True)
        line(slide, x + 0.28, 4.02, x + 3.20, 4.04, SLATE)
        text_box(slide, "STRENGTH", x + 0.28, 4.32, 1.0, 0.22, 8, GREEN, True)
        text_box(slide, strength, x + 0.28, 4.62, 2.92, 0.45, 11.5, WHITE, True)
        text_box(slide, "TRADE-OFF", x + 0.28, 5.30, 1.0, 0.22, 8, ORANGE, True)
        text_box(slide, tradeoff, x + 0.28, 5.60, 2.92, 0.42, 11.5, WHITE)
    text_box(slide, "The score scales differ; compare rankings and metrics—not raw scores.", 0.78, 6.65, 11.05, 0.28, 11, MUTED, True, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # 7 — Evaluation
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 7, "Evaluation", "Complementary metrics answer different questions",
               source="Evaluation definitions: Introduction to Information Retrieval, Chapter 8 · nDCG: Järvelin & Kekäläinen (2002)")
    metrics = [
        ("P@k", "How pure is the top-k list?", BLUE),
        ("R@k", "How much relevant material was found?", CYAN),
        ("MAP", "Are all relevant items ranked early?", MAGENTA),
        ("MRR", "How soon is the first useful result?", ORANGE),
        ("nDCG@k", "Is relevance concentrated near the top?", GREEN),
    ]
    for index, (name, question, color) in enumerate(metrics):
        row = index // 3
        column = index % 3
        x = 0.78 + column * 4.0
        y = 1.50 + row * 1.45
        w = 3.70 if index < 3 else 5.75
        if index >= 3:
            x = 0.78 + (index - 3) * 5.98
        rect(slide, x, y, w, 1.12, WHITE, True, line=LINE)
        rect(slide, x + 0.18, y + 0.19, 0.92, 0.72, color, True)
        text_box(slide, name, x + 0.18, y + 0.19, 0.92, 0.72, 14, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        text_box(slide, question, x + 1.28, y + 0.22, w - 1.52, 0.64, 12, INK, True, valign=MSO_ANCHOR.MIDDLE)
    rect(slide, 0.78, 4.65, 11.75, 1.52, NAVY, True)
    text_box(slide, "Evaluation protocol", 1.05, 4.94, 1.80, 0.30, 12, ORANGE, True)
    protocol = ["8 judged queries", "Full ranking for MAP / MRR", "k = 1, 3, 5, 10", "500 timing repetitions"]
    for index, value in enumerate(protocol):
        x = 3.05 + index * 2.25
        rect(slide, x, 4.89, 1.88, 0.64, NAVY_2, True, line=SLATE)
        text_box(slide, value, x + 0.10, 4.89, 1.68, 0.64, 10, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Effectiveness is the primary outcome. Sub-millisecond latency is descriptive only.", 1.05, 5.65, 10.90, 0.26, 9.5, MUTED, False, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # 8 — Results
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 8, "Evaluation results", "Ranking quality differs even when recall is identical")
    add_metric(slide, 0.78, 1.39, 2.22, 1.26, "0.325", "P@5 · ALL MODELS", BLUE)
    add_metric(slide, 3.17, 1.39, 2.22, 1.26, "1.000", "R@5 · ALL MODELS", CYAN)
    rect(slide, 5.73, 1.39, 6.80, 1.26, "EAF0FF", True, line="C8D6FF")
    text_box(slide, "Same coverage, different ordering", 6.05, 1.62, 5.95, 0.30, 15, INK, True, FONT_DISPLAY)
    text_box(slide, "MAP, MRR, and nDCG reveal differences hidden by P@5 and R@5.", 6.05, 2.03, 5.95, 0.28, 10.5, "52637A")
    for y, (model, result), color in zip(
        (3.16, 4.20, 5.24),
        METRICS["models"].items(),
        (BLUE, MAGENTA, CYAN),
    ):
        add_model_result_row(slide, y, model, color, result["mean"])
    rect(slide, 0.78, 6.20, 11.75, 0.56, NAVY, True)
    text_box(slide, "Winner depends on the question", 1.05, 6.31, 2.35, 0.26, 10, ORANGE, True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "TF–IDF leads MAP · Binary overlap leads MRR and nDCG@5 · BM25 remains competitive", 3.25, 6.29, 8.65, 0.30, 11, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    set_slide_transition(slide)

    # 9 — Major findings
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 9, "Major findings", "What changed across models and cut-offs")
    add_card(slide, 0.78, 1.48, 4.10, 1.35, "TF–IDF: best complete ranking", "MAP = 0.948. It ranked both relevant items ideally for Q3 and Q4.", MAGENTA, number="01")
    add_card(slide, 0.78, 3.03, 4.10, 1.35, "Overlap: best first answer", "MRR = 1.000 and P@1 = 1.000, aided by simple term coverage and deterministic ties.", BLUE, number="02")
    add_card(slide, 0.78, 4.58, 4.10, 1.35, "Differences are concentrated", "Five of eight queries tie at AP = 1.000. Q3, Q4, and Q8 drive the model ordering.", ORANGE, number="03")
    rect(slide, 5.25, 1.48, 7.28, 4.80, WHITE, True, line=LINE)
    slide.shapes.add_picture(str(WEEK7 / "screenshots/recall_cutoffs.png"), Inches(5.52), Inches(1.80), width=Inches(6.72))
    text_box(slide, "Recall reaches 1.000 by rank 3 for every model; larger k adds non-relevant items and lowers precision.",
             5.72, 5.80, 6.35, 0.38, 10.5, "52637A", True, align=PP_ALIGN.CENTER)
    rect(slide, 0.78, 6.35, 11.75, 0.42, "FFF3DF", True, line="F6D69A")
    text_box(slide, "Do not read a 0.010 MAP difference as decisive: eight queries are too few for a stable winner.",
             1.02, 6.43, 11.20, 0.22, 10, "8A4B08", True, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # 10 — Sample output
    slide = presentation.slides.add_slide(blank)
    add_background(slide, NAVY)
    add_header(slide, 10, "Sample output", "One query exposes three ranking behaviors", dark=True)
    rect(slide, 0.72, 1.42, 11.90, 4.95, "0D131A", True, line=SLATE)
    add_picture_cover(slide, WEEK7 / "screenshots/ranking_output.png", 0.86, 1.58, 11.62, 4.63)
    rect(slide, 0.78, 6.48, 11.75, 0.58, NAVY_2, True, line=SLATE)
    text_box(slide, "Q8 insight", 1.02, 6.62, 1.05, 0.26, 9.5, ORANGE, True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Overlap ranks relevant S012 first; TF–IDF and BM25 rank non-relevant S003 first because several weighted terms match.",
             2.07, 6.57, 9.90, 0.32, 10.5, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    set_slide_transition(slide)

    # 11 — Critical analysis
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, 11, "Critical analysis", "Strong implementation evidence; limited model-selection evidence")
    rect(slide, 0.78, 1.48, 5.72, 4.95, "EBF8F2", True, line="BDE8D3")
    rect(slide, 6.78, 1.48, 5.75, 4.95, "FFF1F1", True, line="F4C7C9")
    text_box(slide, "What is strong", 1.10, 1.84, 2.4, 0.40, 20, "155E3E", True, FONT_DISPLAY)
    bullet_list(slide, [
        "Identical controls across models",
        "Transparent, dependency-free scoring",
        "Known-answer metric tests pass",
        "Aggregate and per-query evidence",
        "Reproducible JSON, code, and figures",
    ], 1.10, 2.55, 4.85, 3.35, 14, "245844", GREEN, 11)
    text_box(slide, "What limits the claim", 7.10, 1.84, 3.1, 0.40, 20, "8B2F35", True, FONT_DISPLAY)
    bullet_list(slide, [
        "Only 18 documents and 8 queries",
        "Manual—not official—judgments",
        "Tie order benefits the overlap model",
        "BM25 parameters were not tuned",
        "Latency is too small to generalize",
    ], 7.10, 2.55, 4.85, 3.35, 14, "74343A", RED, 11)
    rect(slide, 0.78, 6.58, 11.75, 0.44, NAVY, True)
    text_box(slide, "Claim boundary  →  the evaluation proves the pipeline, not the best model for full CISI.",
             1.00, 6.67, 11.30, 0.22, 10.5, WHITE, True, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # 12 — Recommendations and conclusion
    slide = presentation.slides.add_slide(blank)
    add_background(slide, NAVY)
    add_header(slide, 12, "Recommendations", "Move from demonstration to defensible model selection", dark=True,
               source="Project repositories prepared: github.com/Chief-spartan-117/week-6-ir-system · /week-7-ir-evaluation")
    recommendations = [
        ("1", "Run full CISI", "Evaluate all 1,460 documents and official judgments.", BLUE),
        ("2", "Tune fairly", "Use development queries for BM25 k1/b; keep test queries untouched.", CYAN),
        ("3", "Quantify uncertainty", "Report paired query differences and confidence intervals.", ORANGE),
        ("4", "Address vocabulary gaps", "Compare stronger stemming and controlled query expansion.", MAGENTA),
    ]
    for index, (number, title, body, color) in enumerate(recommendations):
        x = 0.78 + (index % 2) * 6.02
        y = 1.48 + (index // 2) * 1.48
        add_card(slide, x, y, 5.72, 1.20, title, body, color, dark=True, number=number)
    rect(slide, 0.78, 4.68, 11.75, 1.52, "202B48", True, line=SLATE)
    text_box(slide, "Recommended baseline", 1.10, 4.98, 2.40, 0.30, 10, CYAN, True)
    text_box(slide, "Retain TF–IDF now.", 1.10, 5.37, 3.30, 0.46, 23, WHITE, True, FONT_DISPLAY)
    text_box(slide, "It delivered the best demonstration MAP, remains explainable, and requires no fitted parameters. Reconsider only after full-collection evidence.",
             4.25, 5.05, 7.55, 0.78, 13, "C9D4E7", False, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Build → compare → measure → challenge the evidence → iterate", 0.78, 6.62, 11.75, 0.32, 12, ORANGE, True, align=PP_ALIGN.CENTER)
    set_slide_transition(slide)

    # Document metadata
    presentation.core_properties.title = "Evaluation of Information Retrieval Models"
    presentation.core_properties.subject = "Week 8 summary of Week 6 and Week 7 CISI IR project"
    presentation.core_properties.author = "Sagar Shrestha"
    presentation.core_properties.keywords = "information retrieval, CISI, TF-IDF, BM25, evaluation, MAP, nDCG"
    presentation.core_properties.comments = "Generated from executed Week 6–7 project evidence."
    return presentation


if __name__ == "__main__":
    deck = build_deck()
    deck.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(deck.slides)} slides)")

