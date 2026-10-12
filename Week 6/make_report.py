"""Create the submission-ready DOCX report for the Week 6 IR assignment."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUT = BASE / "report.docx"
FIGURES = BASE / "screenshots"
METRICS = json.loads((FIGURES / "metrics.json").read_text(encoding="utf-8"))
GITHUB_URL = "https://github.com/Chief-spartan-117/week-6-ir-system"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    tc_pr.append(fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_apa_reference(doc: Document, segments: list[tuple[str, bool]], url: str | None = None) -> None:
    """Add an APA 7 reference entry with a 0.5-inch hanging indent."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.space_after = Pt(8)
    for value, italic in segments:
        run = paragraph.add_run(value)
        run.italic = italic
    if url:
        add_hyperlink(paragraph, url, url)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)


def add_figure(doc: Document, filename: str, caption: str, width: float = 6.7) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    add_caption(doc, caption)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Code")
    paragraph.add_run(text)


def add_metric_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headings = ["Query", "P@5", "R@5", "AP", "RR", "nDCG@5"]
    for cell, heading in zip(table.rows[0].cells, headings):
        shade(cell, "1E3A5F")
        set_cell_text(cell, heading, bold=True, color="FFFFFF")
    for row in METRICS["rows"]:
        cells = table.add_row().cells
        values = [
            row["query_id"], f'{row["precision"]:.3f}', f'{row["recall"]:.3f}',
            f'{row["ap"]:.3f}', f'{row["rr"]:.3f}', f'{row["ndcg"]:.3f}',
        ]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
    cells = table.add_row().cells
    mean = METRICS["mean"]
    values = ["Mean", f'{mean["precision"]:.3f}', f'{mean["recall"]:.3f}', f'{mean["ap"]:.3f}', f'{mean["rr"]:.3f}', f'{mean["ndcg"]:.3f}']
    for cell, value in zip(cells, values):
        shade(cell, "E2E8F0")
        set_cell_text(cell, value, bold=True)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13
    for name, size, color in (("Title", 28, "16324F"), ("Heading 1", 18, "16324F"), ("Heading 2", 13, "2563EB"), ("Heading 3", 11, "334155")):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string("475569")

    if "Code" not in styles:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code"]
    code.font.name = "Menlo"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("E2E8F0")
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code_element = code.element
    paragraph_properties = code_element.get_or_add_pPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), "172033")
    paragraph_properties.append(fill)

    header = section.header.paragraphs[0]
    header.text = "WEEK 6  /  INFORMATION RETRIEVAL SYSTEM"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    add_page_number(section.footer.paragraphs[0])


def title_page(doc: Document) -> None:
    doc.add_paragraph("WEEK 6 ASSIGNMENT", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\n")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Design and Implementation of a\nBasic Information Retrieval System")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("TF–IDF Vector Space Retrieval on the CISI Collection")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string("2563EB")
    doc.add_paragraph("\n")

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    details = [
        ("Student", "Sagar Shrestha"),
        ("Student ID", "____________________________"),
        ("Module / Lecturer", "____________________________"),
        ("Submission date", date.today().strftime("%d %B %Y")),
        ("Programming language", "Python 3"),
    ]
    for row, (label, value) in zip(table.rows, details):
        shade(row.cells[0], "E2E8F0")
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Repository: ").bold = True
    add_hyperlink(p, GITHUB_URL, GITHUB_URL)
    note = doc.add_paragraph("Repository prepared locally; remote publication requires GitHub authentication before submission.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].font.size = Pt(9)
    note.runs[0].italic = True
    note.runs[0].font.color.rgb = RGBColor.from_string("9A3412")
    doc.add_page_break()


def build_report() -> Document:
    doc = Document()
    configure_document(doc)
    title_page(doc)

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This project proposes and implements a basic information retrieval system for the public CISI test collection. "
        "The system parses structured CISI records, applies a consistent text-processing pipeline, constructs an inverted "
        "index, and ranks documents with logarithmic TF–IDF weights and cosine similarity. It also includes query parsing, "
        "relevance-judgment parsing, five evaluation measures, a command-line interface, and automated tests."
    )
    doc.add_paragraph(
        "Because this execution environment cannot directly download external files, the repository contains an 18-record "
        "offline demonstration corpus derived from CISI subject areas and eight manually judged queries. It is used only to "
        "prove the end-to-end implementation. The supplied downloader and evaluation command are designed for the full CISI "
        "files. Demonstration results were P@5 = 0.325, R@5 = 1.000, MAP = 0.948, MRR = 0.938, and nDCG@5 = 0.962. "
        "These deliberately small-sample results are not claimed as full CISI benchmark scores."
    )

    doc.add_heading("Contents", level=1)
    contents = [
        "1. Introduction and objectives", "2. Dataset", "3. System architecture and design",
        "4. Implementation", "5. Experimental method", "6. Results and discussion",
        "7. Verification", "8. Limitations and future work", "9. Reproducibility and GitHub",
        "10. Conclusion", "References", "Appendix A — Project structure",
    ]
    for item in contents:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    doc.add_heading("1. Introduction and Objectives", level=1)
    doc.add_paragraph(
        "Information retrieval differs from exact database lookup because a natural-language query may only partially match a "
        "document. The system must estimate relevance and return a ranked list. The assignment goal was therefore translated "
        "into an auditable pipeline in which every stage—from raw text to score—can be inspected and explained."
    )
    doc.add_heading("1.1 Objectives", level=2)
    add_bullets(doc, [
        "Use a publicly available IR collection with documents, test queries, and relevance assessments.",
        "Apply identical normalization to documents and queries to avoid representation mismatch.",
        "Build an inverted index rather than scan every full document for every query.",
        "Implement a classical vector-space ranking model without hiding the algorithm behind a library.",
        "Measure ranking quality and distinguish demonstration results from benchmark results.",
        "Package the code, tests, figures, report, and publication instructions as one reproducible project.",
    ])
    doc.add_heading("1.2 Scope", level=2)
    doc.add_paragraph(
        "The implementation is a local, single-process prototype. It supports the CISI format and interactive queries, but it "
        "does not attempt distributed indexing, learning-to-rank, semantic embeddings, user accounts, or a web interface. This "
        "scope keeps the work aligned with a basic IR assignment and makes each design decision visible."
    )

    doc.add_heading("2. Dataset", level=1)
    doc.add_heading("2.1 Full CISI collection", level=2)
    doc.add_paragraph(
        "CISI is a classic information retrieval test collection concerning library and information science. The University "
        "of Glasgow collection page lists 1,460 documents, 112 queries, and a size of about 2.2 MB. It distributes document, "
        "query, relevance-assessment, and Boolean-query files. The collection is suitable for a teaching project because it is "
        "small enough for an in-memory index while still supporting standard relevance-based evaluation (University of Glasgow "
        "Information Retrieval Group, n.d.-a, n.d.-b)."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["File", "Purpose", "Parser behavior"]):
        shade(cell, "1E3A5F"); set_cell_text(cell, text, True, "FFFFFF")
    for values in [
        ("CISI.ALL", "1,460 document records", "Reads .I, .T, .A, and .W; ignores .X links"),
        ("CISI.QRY", "112 natural-language queries", "Splits records at .I markers"),
        ("CISI.REL", "Query–document judgments", "Maps query IDs to sets of relevant IDs"),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values): set_cell_text(cell, value)
    doc.add_heading("2.2 Offline demonstration data", level=2)
    doc.add_paragraph(
        "For guaranteed offline execution, the repository includes 18 concise records representing topics visible in early CISI "
        "records, plus eight demonstration queries and binary relevance sets. The records are stored as JSON. Two are explicitly "
        "marked as demonstration records. This miniature corpus validates program behavior, screenshots, and metric calculations; "
        "it is not a substitute for publishing full-collection benchmark results."
    )
    doc.add_heading("2.3 Data ethics and validity", level=2)
    doc.add_paragraph(
        "The project downloads a public research collection from a named mirror and preserves document identifiers. It does not "
        "collect personal data or infer sensitive attributes. The main validity risk is evaluation on a very small, manually "
        "constructed sample. The report mitigates this by labeling every sample result and providing the exact command required "
        "for the full collection."
    )
    doc.add_page_break()

    doc.add_heading("3. System Architecture and Design", level=1)
    add_figure(doc, "architecture.png", "Figure 1. End-to-end architecture of the implemented retrieval system.")
    doc.add_heading("3.1 Preprocessing", level=2)
    add_numbered(doc, [
        "Unicode NFKD normalization converts accented forms into a stable ASCII representation.",
        "Text is lowercased so capitalization does not create separate vocabulary entries.",
        "A regular expression retains alphabetic tokens and internal hyphens.",
        "A fixed English stop list removes high-frequency grammatical words with little discriminating value.",
        "Conservative suffix rules reduce forms such as ‘libraries’ and ‘indexing’ to shared stems.",
    ])
    doc.add_paragraph(
        "The title is included twice in each document’s searchable text. This provides a small field boost because a concise title "
        "usually expresses the document topic more directly than the abstract. Authors remain searchable but are not specially weighted."
    )
    doc.add_heading("3.2 Index structure", level=2)
    doc.add_paragraph(
        "The dictionary maps each term to a postings dictionary of document ID → raw term frequency. A separate document-frequency "
        "counter supports IDF calculation. Precomputed L2 document norms avoid recalculating full vector lengths for each query. "
        "Only postings for query terms are traversed during scoring, which is more efficient than comparing a query with every token "
        "of every document."
    )
    doc.add_heading("3.3 Retrieval model", level=2)
    doc.add_paragraph(
        "The vector-space design follows established TF–IDF and cosine-scoring principles (Manning et al., 2008). "
        "For term t in document d, the implementation uses:"
    )
    add_code(doc, "tf_weight(t,d) = 1 + ln(tf(t,d))\nidf(t) = ln((N + 1) / (df(t) + 1)) + 1\nw(t,d) = tf_weight(t,d) × idf(t)")
    doc.add_paragraph(
        "Adding one inside the corpus-size and document-frequency terms smooths IDF, keeps it finite, and gives every indexed term "
        "a positive weight. Query terms use the same weighting. The final score is the cosine of the query and document vectors:"
    )
    add_code(doc, "score(q,d) = (q · d) / (||q||₂ × ||d||₂)")
    doc.add_paragraph(
        "Cosine normalization reduces the tendency of long documents to score highly merely because they contain more words. Results "
        "are sorted by descending score and then by document ID for deterministic tie breaking."
    )
    doc.add_heading("3.4 Query processing", level=2)
    doc.add_paragraph(
        "A user query passes through the same preprocessing function as the documents. Terms absent from the vocabulary are ignored. "
        "If no query term exists in the index, the system returns an empty ranked list rather than arbitrary zero-score records. "
        "The top-k parameter controls output length and is validated as a positive integer."
    )

    doc.add_heading("4. Implementation", level=1)
    doc.add_heading("4.1 Main components", level=2)
    components = [
        ("ir_system.py", "Document models, CISI parser, preprocessing, index construction, cosine ranking, CLI"),
        ("evaluate.py", "Query/judgment parsers and P@k, Recall@k, AP, RR, nDCG@k aggregation"),
        ("download_data.py", "Downloads three files and validates 1,460 documents and 112 queries"),
        ("demo.py", "Runs three representative searches on bundled data"),
        ("tests/", "Six regression tests covering tokenization, ranking, validation, persistence, and parsing"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Component", "Responsibility"]):
        shade(cell, "1E3A5F"); set_cell_text(cell, text, True, "FFFFFF")
    for values in components:
        cells = table.add_row().cells
        for cell, value in zip(cells, values): set_cell_text(cell, value)
    doc.add_heading("4.2 Preprocessing implementation snapshot", level=2)
    add_figure(doc, "code_preprocessing.png", "Figure 2. Actual preprocessing function used by both documents and queries.", width=6.5)
    doc.add_paragraph(
        "Figure 2 shows the deliberately short shared pipeline. Keeping preprocessing in one function prevents training/serving skew: "
        "an indexed term and a query term can match only if they undergo identical normalization."
    )
    doc.add_heading("4.3 Error handling and reproducibility", level=2)
    add_bullets(doc, [
        "An empty document collection raises a clear ValueError.",
        "A non-positive top-k value is rejected.",
        "Unknown-vocabulary queries return an empty list.",
        "The downloader validates collection counts before reporting success.",
        "Ranking ties are deterministic, and all sample inputs are versioned as JSON.",
        "Retrieval code has no third-party runtime dependency; python-docx and Pillow are only needed for artifacts.",
    ])
    doc.add_page_break()

    doc.add_heading("5. Experimental Method", level=1)
    doc.add_heading("5.1 Procedure", level=2)
    add_numbered(doc, [
        "Load the 18 demonstration documents and build the in-memory inverted index.",
        "Run each of eight natural-language queries through the shared preprocessing pipeline.",
        "Retrieve a full ranking for AP and RR, while recording cut-off metrics at k = 5.",
        "Compare document IDs with the binary relevance set for each query.",
        "Average each measure across all evaluated queries.",
        "Run the six automated tests and capture the output used in this report.",
    ])
    doc.add_heading("5.2 Measures", level=2)
    doc.add_paragraph(
        "The evaluation measures follow standard ranked-retrieval definitions described by Manning et al. (2008)."
    )
    measures = [
        ("P@5", "Relevant documents in the first five positions divided by five."),
        ("R@5", "Relevant documents found in the first five divided by all judged relevant documents."),
        ("AP / MAP", "Precision measured at every relevant hit; MAP is its mean across queries."),
        ("RR / MRR", "Reciprocal rank of the first relevant result; averaged across queries."),
        ("nDCG@5", "Position-discounted gain normalized by an ideal ordering at cut-off five."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Measure", "Interpretation"]):
        shade(cell, "1E3A5F"); set_cell_text(cell, text, True, "FFFFFF")
    for values in measures:
        cells = table.add_row().cells
        for cell, value in zip(cells, values): set_cell_text(cell, value)
    doc.add_paragraph(
        "P@5 uses a fixed denominator of five even when fewer than five relevant records exist. Therefore, a query with only one "
        "known relevant item has a maximum P@5 of 0.20. This explains why low P@5 can coexist with perfect recall and first-rank performance."
    )

    doc.add_heading("6. Results and Discussion", level=1)
    doc.add_heading("6.1 Query-level effectiveness", level=2)
    add_metric_table(doc)
    add_caption(doc, "Table 1. Offline demonstration results at k = 5. These are not full CISI benchmark scores.")
    add_figure(doc, "metrics_chart.png", "Figure 3. Mean effectiveness on the bundled demonstration queries.", width=6.3)
    doc.add_heading("6.2 Interpretation", level=2)
    doc.add_paragraph(
        "The system retrieved every judged relevant record within the first five positions for all eight queries, so mean R@5 was "
        "1.000. Seven queries placed a relevant item first. Query Q8 was the exception: its first relevant item appeared at rank two, "
        "and its second relevant item appeared later, producing AP = 0.583, RR = 0.500, and nDCG@5 = 0.693. This illustrates that "
        "recall alone does not describe ranking quality."
    )
    doc.add_paragraph(
        "Mean P@5 was 0.325 because most queries have only one or two known relevant documents. Even an ideal ranking cannot fill all "
        "five positions with relevant records in this sample. MAP and nDCG are more informative here because they reward early placement "
        "while respecting the available judgments."
    )
    doc.add_heading("6.3 Example retrieval output", level=2)
    add_figure(doc, "demo_results.png", "Figure 4. Three searches produced by the implemented command-line demonstration.", width=6.55)
    doc.add_paragraph(
        "For ‘automatic text indexing and document retrieval,’ the automatic-indexing record is ranked first, followed by abstracting "
        "and thesaurus records. For the university-library query, systems analysis and administration appear above building design. "
        "The evaluation query strongly favors the record explicitly discussing precision, recall, and ranked retrieval. These examples "
        "show understandable lexical matching rather than opaque model behavior."
    )
    add_figure(doc, "evaluation.png", "Figure 5. Actual evaluation program output used to populate Table 1.", width=6.55)
    doc.add_heading("6.4 Performance characteristics", level=2)
    doc.add_paragraph(
        "Index construction is linear in the number of processed tokens. Query scoring touches postings only for terms present in the "
        "query, then sorts the non-zero candidates. The design is appropriate for CISI’s 1,460 documents. At web scale, the Python "
        "dictionaries and full candidate sort would need compressed postings, skip data, top-k heaps, sharding, and persistent storage."
    )
    doc.add_page_break()

    doc.add_heading("7. Verification", level=1)
    doc.add_paragraph(
        "Six automated unit tests were executed with Python’s built-in unittest framework. They verify stop-word removal and stemming, "
        "ranking of an obviously relevant document, empty results for an out-of-vocabulary query, validation of top-k, valid JSON index "
        "serialization, and correct parsing of CISI record fields. All tests passed."
    )
    add_figure(doc, "tests.png", "Figure 6. Automated verification output: six tests passed.", width=6.55)
    doc.add_heading("7.1 Acceptance check", level=2)
    checklist = [
        ("Dataset documented", "Yes — full CISI files and offline demonstration layer are distinguished."),
        ("Architecture designed", "Yes — preprocessing, index, model, and query path are specified."),
        ("System implemented", "Yes — parser, index, retrieval, CLI, evaluation, and tests run."),
        ("Code/results snapshots", "Yes — Figures 2, 4, 5, and 6 are generated from current files/output."),
        ("DOCX report", "Yes — generated as report.docx."),
        ("GitHub publication", "Pending authentication — repository is prepared locally with the intended URL."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Requirement", "Evidence/status"]):
        shade(cell, "1E3A5F"); set_cell_text(cell, text, True, "FFFFFF")
    for values in checklist:
        cells = table.add_row().cells
        for cell, value in zip(cells, values): set_cell_text(cell, value)

    doc.add_heading("8. Limitations and Future Work", level=1)
    doc.add_heading("8.1 Limitations", level=2)
    add_bullets(doc, [
        "Demonstration judgments are small and manually authored, so the high scores should not be generalized.",
        "Suffix stripping is intentionally simple and may conflate unrelated terms or miss some morphological variants.",
        "Pure lexical matching cannot bridge synonyms when query and document use different vocabulary.",
        "The fixed stop list and title boost were selected heuristically rather than tuned on held-out queries.",
        "The index is in memory and the save method currently supports inspection, not a load-optimized production format.",
    ])
    doc.add_heading("8.2 Future improvements", level=2)
    add_numbered(doc, [
        "Run the supplied full-collection command and report official CISI metrics over all judged queries.",
        "Compare TF–IDF with BM25 using the same preprocessing and judgments.",
        "Use a tested Porter or Snowball stemmer and perform an ablation study for stop-word removal and title boosting.",
        "Add pseudo-relevance feedback or thesaurus-based expansion to reduce vocabulary mismatch.",
        "Create a lightweight web interface and persistent compressed index after the retrieval baseline is validated.",
    ])

    doc.add_heading("9. Reproducibility and GitHub", level=1)
    doc.add_paragraph("Offline reproduction from the Week 6 directory:")
    add_code(doc, "python3 demo.py\npython3 evaluate.py --sample --top-k 5\npython3 -m unittest discover -s tests -v")
    doc.add_paragraph("Full CISI reproduction with internet access:")
    add_code(doc, "python3 download_data.py\npython3 evaluate.py --data-dir data --top-k 10\npython3 ir_system.py --corpus data/CISI.ALL --query \"automatic indexing\" --top-k 5")
    doc.add_paragraph("Publish the prepared local repository:")
    add_code(doc, "gh auth login\ngh repo create Chief-spartan-117/week-6-ir-system --public\ngit push -u origin main")
    p = doc.add_paragraph()
    p.add_run("Intended GitHub link: ").bold = True
    add_hyperlink(p, GITHUB_URL, GITHUB_URL)
    doc.add_paragraph(
        "Before submission, confirm that the link opens in a private browser window. If a different repository name is chosen, update "
        "GITHUB_URL near the top of make_report.py and regenerate this file."
    )

    doc.add_heading("10. Conclusion", level=1)
    doc.add_paragraph(
        "The assignment produced a complete basic retrieval pipeline rather than only a similarity calculation. CISI records can be "
        "downloaded and parsed, document text is normalized consistently, an inverted index supports efficient candidate scoring, and "
        "TF–IDF cosine similarity creates a deterministic ranked result. Evaluation and tests make the behavior measurable and repeatable. "
        "The offline experiment demonstrates that the implementation functions end to end, while explicit labeling prevents its small "
        "sample from being confused with a full benchmark. The natural next experiment is a controlled TF–IDF versus BM25 comparison on "
        "all CISI judgments."
    )

    doc.add_page_break()
    doc.add_heading("References", level=1)
    add_apa_reference(doc, [
        ("GianRomani. (n.d.). ", False),
        ("CISI-project-MLOps", True),
        (" [Data set]. GitHub. Retrieved October 12, 2026, from ", False),
    ], "https://github.com/GianRomani/CISI-project-MLOps")
    add_apa_reference(doc, [
        ("Manning, C. D., Raghavan, P., & Schütze, H. (2008). ", False),
        ("Introduction to information retrieval", True),
        (". Cambridge University Press. ", False),
    ], "https://nlp.stanford.edu/IR-book/")
    add_apa_reference(doc, [
        ("University of Glasgow Information Retrieval Group. (n.d.-a). ", False),
        ("CISI collection", True),
        (". Retrieved October 12, 2026, from ", False),
    ], "https://ir.dcs.gla.ac.uk/resources/test_collections/cisi/")
    add_apa_reference(doc, [
        ("University of Glasgow Information Retrieval Group. (n.d.-b). ", False),
        ("Test collections", True),
        (". Retrieved October 12, 2026, from ", False),
    ], "https://ir.dcs.gla.ac.uk/resources/test_collections/")

    doc.add_heading("Appendix A — Project Structure", level=1)
    add_code(doc, """Week 6/
├── data/
│   ├── sample_documents.json
│   └── sample_queries.json
├── screenshots/
├── tests/test_ir_system.py
├── demo.py
├── download_data.py
├── evaluate.py
├── ir_system.py
├── make_report.py
├── make_screenshots.py
├── README.md
├── requirements.txt
└── report.docx""")
    doc.add_paragraph(
        "Generated artifacts remain reproducible: make_screenshots.py reruns the program and rebuilds all figures, and make_report.py "
        "reads the machine-readable metrics before constructing this document."
    )
    return doc


if __name__ == "__main__":
    report = build_report()
    report.save(OUT)
    print(f"Saved {OUT}")

